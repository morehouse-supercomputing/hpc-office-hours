"""Build the HPC Office Hours deck (Morehouse brand). Copy source: docs/hpc-office-hours-slide-text.md."""
import random
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"
OUT = HERE / "hpc-office-hours.pptx"

MAROON = RGBColor(0x84, 0x00, 0x28)
GOLD = RGBColor(0xC1, 0xA2, 0x31)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xA7, 0xA8, 0xAA)
LIGHT = RGBColor(0xF4, 0xF1, 0xF2)
INK2 = RGBColor(0x4A, 0x4A, 0x4A)
HEAD = "Arial"
BODY = "Georgia"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def network_bg():
    """Faint node-and-edge graph for the title slide."""
    path = IMG / "network-bg.png"
    W, H = 2666, 1500
    random.seed(7)
    im = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    pts = [(random.randint(0, W), random.randint(0, H)) for _ in range(70)]
    for i, (x, y) in enumerate(pts):
        near = sorted(pts, key=lambda p: (p[0] - x) ** 2 + (p[1] - y) ** 2)[1:4]
        for p in near:
            d.line([(x, y), p], fill=(193, 162, 49, 55), width=3)
    for x, y in pts:
        r = random.choice([6, 9, 12])
        d.ellipse([x - r, y - r, x + r, y + r], fill=(255, 255, 255, 70))
    im.save(path)
    return path


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def box(slide, x, y, w, h, fill=None, line=None, lw=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = lw or Pt(1)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, runs, size=18, color=BLACK, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, overrides)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for seg in ([(para, {})] if isinstance(para, str) else para):
            t, o = seg if isinstance(seg, tuple) else (seg, {})
            r = p.add_run()
            r.text = t
            r.font.size = Pt(o.get("size", size))
            r.font.name = o.get("font", font)
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def title(slide, t, color=MAROON, y=Inches(0.55)):
    text(slide, Inches(0.7), y, Inches(11.9), Inches(0.9), t.upper(), size=34,
         color=color, font=HEAD, bold=True)
    box(slide, Inches(0.7), y + Inches(0.92), Inches(1.1), Inches(0.08), fill=GOLD)


def footer(slide, dark=False):
    c = GRAY if not dark else RGBColor(0xD8, 0xB8, 0xC2)
    text(slide, Inches(0.7), Inches(7.0), Inches(8), Inches(0.3),
         "MOREHOUSE SUPERCOMPUTING FACILITY  ·  HPC OFFICE HOURS", size=10,
         color=c, font=HEAD, bold=True)


def picture(slide, path, x, y, w, h):
    """Place an image cropped to fill the box."""
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(str(path), x, y, w, h)
    tr, ir = w / h, iw / ih
    if ir > tr:
        c = (1 - tr / ir) / 2
        pic.crop_left = pic.crop_right = c
    else:
        c = (1 - ir / tr) / 2
        pic.crop_top = c * 0.6
        pic.crop_bottom = c * 1.4
    return pic


def highlight(slide, y, label, body):
    box(slide, Inches(0.7), y, Inches(11.93), Inches(0.95), fill=MAROON)
    box(slide, Inches(0.7), y, Inches(0.12), Inches(0.95), fill=GOLD)
    text(slide, Inches(1.1), y, Inches(11.3), Inches(0.95),
         [[(label + "  ", {"font": HEAD, "bold": True, "color": GOLD, "size": 16}),
           (body, {"color": WHITE, "size": 20})]], anchor=MSO_ANCHOR.MIDDLE)


def cards(slide, items, y, h, cols=None, numbered=True, head_size=20, body_size=16):
    cols = cols or len(items)
    gap = Inches(0.3)
    w = int((Inches(11.93) - gap * (cols - 1)) / cols)
    for i, (head, body) in enumerate(items):
        x = Inches(0.7) + (w + gap) * i
        box(slide, x, y, w, h, fill=LIGHT)
        box(slide, x, y, w, Inches(0.08), fill=MAROON)
        top = y + Inches(0.3)
        if numbered:
            text(slide, x + Inches(0.3), top, Inches(1), Inches(0.7), f"{i + 1:02d}",
                 size=30, color=GOLD, font=HEAD, bold=True)
            top += Inches(0.75)
        text(slide, x + Inches(0.3), top, w - Inches(0.6), Inches(0.9), head,
             size=head_size, color=MAROON, font=HEAD, bold=True)
        if body:
            text(slide, x + Inches(0.3), top + Inches(0.75), w - Inches(0.6),
                 h - (top - y) - Inches(0.9), body, size=body_size, color=INK2)


def table(slide, x, y, colw, rowh, header, rows, first_col_label=False):
    tbl = slide.shapes.add_table(len(rows) + 1, len(header), x, y,
                                 sum(colw, Inches(0)), rowh * (len(rows) + 1)).table
    for j, w in enumerate(colw):
        tbl.columns[j].width = w
    for i in range(len(rows) + 1):
        tbl.rows[i].height = rowh
        for j in range(len(header)):
            cell = tbl.cell(i, j)
            val = header[j] if i == 0 else rows[i - 1][j]
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = MAROON if (j or not first_col_label) else WHITE
            elif first_col_label and j == 0:
                cell.fill.fore_color.rgb = LIGHT
            else:
                cell.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            tf = cell.text_frame
            tf.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.2)
            p = tf.paragraphs[0]
            r = p.runs[0] if p.runs else p.add_run()
            head = i == 0 or (first_col_label and j == 0)
            r.font.name = HEAD if head else BODY
            r.font.bold = head
            r.font.size = Pt(18 if i == 0 else 16)
            r.font.color.rgb = WHITE if i == 0 else (MAROON if head else BLACK)


# 1 · Welcome
s = prs.slides.add_slide(BLANK)
bg(s, MAROON)
s.shapes.add_picture(str(network_bg()), 0, 0, SW, SH)
box(s, Inches(0.7), Inches(2.35), Inches(0.12), Inches(2.3), fill=GOLD)
text(s, Inches(1.1), Inches(2.2), Inches(11), Inches(0.5), "MOREHOUSE SUPERCOMPUTING FACILITY (MSF)",
     size=16, color=GOLD, font=HEAD, bold=True)
text(s, Inches(1.1), Inches(2.75), Inches(11.5), Inches(1.6), "WELCOME TO\nHPC OFFICE HOURS",
     size=54, color=WHITE, font=HEAD, bold=True, spacing=0.95)
text(s, Inches(1.1), Inches(5.0), Inches(11), Inches(1),
     [[("Dr. Ashley Scruse", {"bold": True}), ("  ·  Deputy Director, MSF", {})],
      "Tonight: getting you onto the supercomputer"], size=20, color=WHITE)

# 2 · Tonight
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Tonight, We Dip Our Toes In")
cards(s, [("Get access", ""), ("Finish onboarding", ""), ("Explore the system", "")],
      Inches(1.9), Inches(2.3), head_size=24)
highlight(s, Inches(4.55), "THE ONE REQUIREMENT",
          "Leave tonight with MFA set up and able to log in.")
text(s, Inches(0.7), Inches(5.85), Inches(11.9), Inches(0.8),
     "Everything else is optional. Follow along or just watch. No jobs, no code, nothing to memorize.",
     size=18, color=INK2)
footer(s)

# 3 · Access
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
picture(s, IMG / "scruse-supercomputer.jpg", Inches(8.4), 0, Inches(4.933), SH)
text(s, Inches(0.7), Inches(0.55), Inches(7.2), Inches(1.9),
     "YOU NOW HAVE ACCESS TO A NATIONAL SUPERCOMPUTER", size=34, color=MAROON, font=HEAD, bold=True)
box(s, Inches(0.7), Inches(2.5), Inches(1.1), Inches(0.08), fill=GOLD)
text(s, Inches(0.7), Inches(3.0), Inches(7.1), Inches(3.5), [
    [("A national academic computing resource, ", {"bold": True}), ("shared by researchers across the country", {})],
    "",
    [("Thousands of computers, called ", {}), ("nodes", {"bold": True, "color": MAROON}), (", working together", {})],
    "",
    "Built for work too big or too slow for a laptop",
], size=22)
footer(s)

# 4 · Systems
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "One Resource, Several Systems")
gap = Inches(0.3)
w = int((Inches(11.93) - gap * 2) / 3)
for i, (name, note, img, focus) in enumerate([
        ("VISTA", "Our focus tonight", "racks-aisle.jpg", True),
        ("STAMPEDE3", "Also available to you", "stampede.jpg", False),
        ("FRONTERA", "Also available to you", "racks-close.jpg", False)]):
    x = Inches(0.7) + (w + gap) * i
    picture(s, IMG / img, x, Inches(1.85), w, Inches(3.1))
    box(s, x, Inches(4.95), w, Inches(1.2), fill=MAROON if focus else LIGHT)
    text(s, x + Inches(0.3), Inches(5.08), w, Inches(0.5), name, size=24, font=HEAD, bold=True,
         color=WHITE if focus else MAROON)
    text(s, x + Inches(0.3), Inches(5.6), w, Inches(0.4), note, size=16,
         color=GOLD if focus else INK2, bold=focus)
text(s, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.5),
     "Same account, same login. Each system is built for different kinds of work.", size=18, color=INK2)
footer(s)

# 5 · When to use it
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "When Do You Need It?")
cards(s, [
    ("Too big", "Your data is too big to open on a laptop"),
    ("Too slow", "Your analysis takes hours or days to run"),
    ("Too many", "You need to run the same thing many times"),
    ("Build skills", "Linux, shared computing, and the tools research runs on"),
], Inches(1.9), Inches(3.2), head_size=22, body_size=16)
highlight(s, Inches(5.45), "GOOD NEWS",
          "Your project doesn't have to need a supercomputer. Learning how to use one counts.")
footer(s)

# 6 · Section divider
s = prs.slides.add_slide(BLANK)
bg(s, MAROON)
picture(s, IMG / "racks-close.jpg", Inches(8.4), 0, Inches(4.933), SH)
box(s, Inches(0.7), Inches(2.6), Inches(0.12), Inches(2.1), fill=GOLD)
text(s, Inches(1.1), Inches(2.45), Inches(7), Inches(0.5), "SECTION", size=16, color=GOLD, font=HEAD, bold=True)
text(s, Inches(1.1), Inches(2.95), Inches(7), Inches(1), "A SHARED SYSTEM", size=50, color=WHITE, font=HEAD, bold=True)
text(s, Inches(1.1), Inches(4.0), Inches(6.8), Inches(1.2),
     "You share this machine with researchers around the world. Shared systems come with rules.",
     size=20, color=WHITE)
footer(s, dark=True)

# 6b · Node vs. core
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Node vs. Core")
gap = Inches(0.3)
w = int((Inches(11.93) - gap * 2) / 3)
for i, (name, what, num, unit) in enumerate([
        ("CORE", "One worker inside a processor", "1", "worker"),
        ("NODE", "One whole computer, full of cores", "144", "cores per CPU node (72 per GPU node)"),
        ("VISTA", "All the nodes wired together", "856", "nodes: 256 CPU + 600 GPU")]):
    x = Inches(0.7) + (w + gap) * i
    box(s, x, Inches(1.9), w, Inches(3.2), fill=MAROON if i == 1 else LIGHT)
    fg, sub = (WHITE, GOLD) if i == 1 else (MAROON, INK2)
    text(s, x + Inches(0.3), Inches(2.15), w - Inches(0.6), Inches(0.5), name, size=22, color=fg, font=HEAD, bold=True)
    text(s, x + Inches(0.3), Inches(2.7), w - Inches(0.6), Inches(0.8), what, size=17, color=WHITE if i == 1 else INK2)
    text(s, x + Inches(0.3), Inches(3.5), w - Inches(0.6), Inches(0.9), num, size=48, color=GOLD, font=HEAD, bold=True)
    text(s, x + Inches(0.3), Inches(4.45), w - Inches(0.6), Inches(0.5), unit, size=15, color=sub)
    if i < 2:
        text(s, x + w, Inches(3.2), gap, Inches(0.6), "›", size=30, color=GOLD, font=HEAD, bold=True,
             align=PP_ALIGN.CENTER)
highlight(s, Inches(5.45), "VISTA IN TOTAL",
          "856 nodes and 80,064 cores working together.")
footer(s)

# 7 · Nodes
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Login Nodes vs. Compute Nodes")
half = Inches(5.9)
for i, (name, rows) in enumerate([
        ("LOGIN NODE", ["Where you land when you log in", "Shared by everyone at once", "Organize files, edit, submit"]),
        ("COMPUTE NODE", ["Where the real work runs", "Reserved for you", "Heavy analysis, big data"])]):
    x = Inches(0.7) + (half + Inches(0.13)) * i
    box(s, x, Inches(1.85), half, Inches(0.75), fill=MAROON if i else BLACK)
    text(s, x, Inches(1.85), half, Inches(0.75), name, size=22, color=WHITE, font=HEAD, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for r, line in enumerate(rows):
        y = Inches(2.6) + Inches(0.78) * r
        box(s, x, y, half, Inches(0.78), fill=LIGHT if r % 2 == 0 else WHITE)
        text(s, x, y, half, Inches(0.78), line, size=19, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
box(s, Inches(6.6), Inches(1.85), Inches(0.13), Inches(3.09), fill=GOLD)
highlight(s, Inches(5.35), "THINK OF IT AS",
          "The login node is the lobby. The compute nodes are the labs.")
footer(s)

# 8 · Files
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Three Places for Your Files")
table(s, Inches(0.7), Inches(1.85), [Inches(2.9)] + [Inches(3.01)] * 3, Inches(0.62),
      ["", "$HOME", "$WORK", "$SCRATCH"], [
          ["Space", "23 GB", "1 TB", "No limit"],
          ["For", "Settings, small scripts", "Your projects", "Big, temporary runs"],
          ["Backed up", "Yes", "No", "No"],
          ["Deleted automatically", "Never", "Never", "After 10 days untouched"],
      ], first_col_label=True)
highlight(s, Inches(5.35), "RULE OF THUMB",
          "Keep your projects in $WORK. Treat $SCRATCH as temporary.")
footer(s)

# 9 · Getting in line
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Three Ways to Get a Compute Node")
cards(s, [
    ("Queue", "Submit a job and wait your turn. First come, first served. Best for long or overnight runs."),
    ("Interactive session", "Work on a compute node live. Good for testing code. Coming in the next workshop."),
    ("Reservation", "Book nodes ahead of time with a ticket in the user portal. \"Two Grace Hopper GPU nodes on Vista, Thursday 8am to 8pm.\""),
], Inches(1.9), Inches(3.2), head_size=22, body_size=16)
highlight(s, Inches(5.45), "TEST BEFORE YOU QUEUE",
          "A script that fails after hours in line sends you to the back of the line.")
footer(s)

# 10 · Rules
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "The Rules of the System")
rules = [
    ("No heavy work on the login node.", "It is shared. Heavy work goes to compute nodes."),
    ("Your account is yours.", "Never share your password or MFA."),
    ("Keep a copy of anything important.", "$SCRATCH is cleaned out automatically."),
    ("Stuck? Ask.", "Breaking something by guessing is worse than asking."),
]
for i, (h, b) in enumerate(rules):
    y = Inches(1.9) + Inches(1.2) * i
    box(s, Inches(0.7), y, Inches(1.0), Inches(1.0), fill=MAROON)
    text(s, Inches(0.7), y, Inches(1.0), Inches(1.0), str(i + 1), size=32, color=GOLD, font=HEAD,
         bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(2.0), y, Inches(10.6), Inches(1.0),
         [[(h + "  ", {"font": HEAD, "bold": True, "color": MAROON, "size": 22}), (b, {"size": 20})]],
         anchor=MSO_ANCHOR.MIDDLE)
footer(s)

# 11 · What's next
s = prs.slides.add_slide(BLANK)
bg(s, MAROON)
text(s, Inches(0.7), Inches(0.45), Inches(11.9), Inches(0.8), "WHAT'S NEXT: REAL DATA FOR YOUR OPIL WORK", size=30, color=WHITE, font=HEAD, bold=True)
box(s, Inches(0.7), Inches(1.25), Inches(1.1), Inches(0.08), fill=GOLD)
text(s, Inches(0.7), Inches(1.5), Inches(11.9), Inches(0.6),
     [[("The data: ", {"bold": True, "color": GOLD}),
       ("the CFPB Consumer Complaint Database. About 10 million real complaints about money transfers, prepaid cards, and digital payments, too big for a laptop.", {})]],
     size=17, color=WHITE)
for i, (date, name, do, leave) in enumerate([
        ("OCT 14", "What the Data Says",
         "Count complaints by product, company, state, and year. Each team edits one cell for its own sector and state.",
         "Three numbers for your Monetization Canvas (due Oct 23)."),
        ("NOV 4", "Finding the Pain Point",
         "Read the complaint narratives at scale and pull the theme closest to your team's business problem.",
         "The pain point, with numbers, for your business model and your Dec 4 pitch.")]):
    x = Inches(0.7) + Inches(6.1) * i
    box(s, x, Inches(2.35), Inches(5.83), Inches(3.05), fill=WHITE)
    box(s, x, Inches(2.35), Inches(0.12), Inches(3.05), fill=GOLD)
    text(s, x + Inches(0.4), Inches(2.5), Inches(5.2), Inches(0.5),
         [[(date + "   ", {"color": MAROON}), (name.upper(), {"color": BLACK})]], size=20, font=HEAD, bold=True)
    text(s, x + Inches(0.4), Inches(3.1), Inches(5.2), Inches(1.2), do, size=15, color=INK2)
    text(s, x + Inches(0.4), Inches(4.3), Inches(5.2), Inches(1.0),
         [[("You leave with: ", {"bold": True, "color": MAROON}), (leave, {})]], size=15, color=BLACK)
box(s, Inches(0.7), Inches(5.65), Inches(11.93), Inches(1.0), fill=WHITE)
box(s, Inches(0.7), Inches(5.65), Inches(0.12), Inches(1.0), fill=GOLD)
text(s, Inches(1.1), Inches(5.65), Inches(11.3), Inches(1.0),
     [[("WHY IT MATTERS  ", {"font": HEAD, "bold": True, "color": MAROON, "size": 15}),
       ("Swap \"we think people struggle with payments\" for \"this many people said so.\" The pain points you find this fall become the requirements for your open payments prototype in the spring.", {"size": 16})]],
     color=BLACK, anchor=MSO_ANCHOR.MIDDLE)
footer(s, dark=True)

# 12 · Let's connect
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Let's Connect")
gap = Inches(0.3)
w = int((Inches(11.93) - gap) / 2)
for i, (head, rows) in enumerate([
        ("DR. ASHLEY SCRUSE", [("LinkedIn", "linkedin.com/in/ashleyscruse"),
                               ("Email", "ashley.scruse@morehouse.edu"),
                               ("GitHub", "github.com/ashleyscruse")]),
        ("STAY LOCKED IN", [("MSF LinkedIn", "[MSF LinkedIn URL]"),
                            ("CBPC LinkedIn", "linkedin.com/company/morehousecbpc"),
                            ("CBPC", "bpccenter.org"),
                            ("MSF GitHub", "github.com/morehouse-supercomputing")])]):
    x = Inches(0.7) + (w + gap) * i
    box(s, x, Inches(1.9), w, Inches(4.4), fill=MAROON if i == 0 else LIGHT)
    box(s, x, Inches(1.9), w, Inches(0.08), fill=GOLD)
    fg = WHITE if i == 0 else MAROON
    text(s, x + Inches(0.4), Inches(2.2), w - Inches(0.8), Inches(0.5), head, size=20, color=GOLD if i == 0 else MAROON, font=HEAD, bold=True)
    for r, (label, val) in enumerate(rows):
        y = Inches(2.95) + Inches(0.8) * r
        text(s, x + Inches(0.4), y, w - Inches(0.8), Inches(0.3), label.upper(), size=12,
             color=RGBColor(0xD8, 0xB8, 0xC2) if i == 0 else INK2, font=HEAD, bold=True)
        text(s, x + Inches(0.4), y + Inches(0.3), w - Inches(0.8), Inches(0.4), val, size=18,
             color=WHITE if i == 0 else BLACK)
text(s, Inches(0.7), Inches(6.45), Inches(11.9), Inches(0.4),
     [[("Guide from tonight: ", {"bold": True, "color": MAROON}), ("morehouse-supercomputing.github.io/hpc-office-hours", {})]],
     size=16, color=BLACK)
footer(s)

NOTES = [
    "Name, Deputy Director, MSF.\nRelaxed hour. No jobs, no code.",
    "Three steps.\nOnly requirement: MFA + login.\nThen: switch to guide page for MFA.",
    "Stragglers keep working on MFA.\nNational resource, shared.\nNodes = computers working together.",
    "Same account, same login.\nVista tonight.",
    "Too big, too slow, too many.\nCard 4: learning counts.",
    "Shared machine = rules.",
    "Core = one worker.\nNode = 144 cores (72 on GPU nodes).\nVista = 856 nodes, 80,064 cores.",
    "Login node = lobby. Shared.\nCompute = labs.\nSets up rule 1.",
    "HOME small, backed up.\nWORK = your projects.\nSCRATCH = temporary, 10 days.\nThese are the folders you'll see in Tapis.",
    "Queue: wait your turn.\nInteractive: next workshop.\nReservation: ticket.\nStory: waited hours, script failed instantly.",
    "Four rules. Read them out.\nThen: switch to Tapis.",
    "CFPB complaints, ~10M.\nOct 14: three numbers for the Canvas.\nNov 4: the pain point.\nFall evidence = spring prototype.",
    "Connect. Guide link.\nQuestions.",
]
for slide, note in zip(prs.slides, NOTES):
    slide.notes_slide.notes_text_frame.text = note

prs.save(OUT)
print(OUT)
